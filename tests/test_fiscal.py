import datetime as dt
import os
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select, update

from caspian.db.models import AuditLog, DocStatus, DocType, Document, Item, Person, PersonKind, StockLedger
from caspian.services import documents as docs
from caspian.services import fiscal, items, master, protected, reports
from caspian.services.backup import SqliteDumper
from caspian.services.errors import ApprovalError, ValidationError
from caspian.services.fiscal import CloseOptions
from caspian.services.fiscal_state import load_state
from caspian.services.items import ItemInput

D = Decimal
PASSWORD = "Backup#Pass2026"
Y1404 = dt.date(2025, 9, 1)  # inside fiscal 1404
OLD = dt.date(2024, 5, 1)  # fiscal 1403, > 12 months before the end of 1404
NEW = dt.date(2026, 4, 10)  # fiscal 1405 (current)


@pytest.fixture
async def env(db, admin, tmp_path):
    u = {x.name: x.id for x in await master.list_units(db)}
    wh = (await master.list_warehouses(db))[0].id
    ids = {}
    for code, name, returnable in (("A", "دریل", False), ("C", "کالای راکد", False),
                                   ("P", "پالت", True), ("N", "بدون گردش", False)):
        ids[code] = await items.create_item(db, admin, ItemInput(code, name, u["عدد"],
                                                                 is_returnable=returnable))
    async with db.session() as s:  # C was defined long ago (creation dates are real clock times)
        await s.execute(update(Item).where(Item.code == "C").values(created_at=dt.datetime(2024, 1, 1)))
    emp = await master.save_person(db, admin, "علی", PersonKind.EMPLOYEE)
    await master.save_person(db, admin, "شخص بی‌استفاده", PersonKind.OTHER)

    async def post(doc_type, date, *lines, **kw):
        return await docs.create_and_post(db, admin, docs.DocumentInput(
            doc_type, date, wh, [docs.LineInput(ids[c], u["عدد"], D(q)) for c, q in lines], **kw))

    await post(DocType.RECEIPT, OLD, ("C", 5))
    await post(DocType.ISSUE, OLD, ("C", 5))  # C: zero stock, last moved in 1403 -> stale
    await post(DocType.RECEIPT, Y1404, ("A", 10), ("P", 6))
    await post(DocType.ISSUE, Y1404, ("A", 3))
    loan = await post(DocType.LOAN_OUT, Y1404, ("P", 4), person_id=emp)  # stays open
    await post(DocType.ISSUE, NEW, ("A", 2))  # already in the new year
    return {"db": db, "admin": admin, "ids": ids, "wh": wh, "u": u, "loan": loan, "emp": emp,
            "dumper": SqliteDumper(db.url.database), "dir": tmp_path / "bk"}


async def _approval(db, admin):
    return await protected.approve(db, admin, protected.ProtectedAction.CLOSE_FISCAL_YEAR, "admin", "4826")


async def _close(env, options=None):
    db, admin = env["db"], env["admin"]
    return await fiscal.run_year_end(db, admin, 1404, options or CloseOptions(purge_stale_items=True),
                                     await _approval(db, admin), env["dumper"], PASSWORD, env["dir"])


async def test_pre_check_and_protection(env):
    db, admin = env["db"], env["admin"]
    check = await fiscal.pre_check(db, admin, 1404)
    assert check.documents == 3 and check.open_loans == 1 and not check.blocking
    assert (await fiscal.pre_check(db, admin, 1405)).not_ended  # current year can't be closed
    with pytest.raises(ApprovalError):
        await fiscal.run_year_end(db, admin, 1404, CloseOptions(), None, env["dumper"], PASSWORD,
                                  env["dir"])
    with pytest.raises(ValidationError):
        await fiscal.run_year_end(db, admin.as_ai(), 1404, CloseOptions(), await _approval(db, admin),
                                  env["dumper"], PASSWORD, env["dir"])
    await docs.create_document(db, admin, docs.DocumentInput(
        DocType.RECEIPT, Y1404, env["wh"], [docs.LineInput(env["ids"]["A"], env["u"]["عدد"], D(1))]))
    assert "پیش‌نویس" in " ".join((await fiscal.pre_check(db, admin, 1404)).blocking)


async def test_close_year_carries_balances_and_loans(env):
    db, admin, ids = env["db"], env["admin"], env["ids"]
    before = {r[0]: r[4] for r in (await reports.stock_balance(db, admin)).rows}
    result = await _close(env)

    assert result.opening_documents == 1 and result.carried_loans == 1
    # C is stale (no stock, last moved in 1403); N is new, just unused, so it stays.
    assert result.removed_items == 1 and result.deactivated_items == 0
    # Current balances unchanged, and the ledger agrees with them (opening + new-year rows).
    after = {r[0]: r[4] for r in (await reports.stock_balance(db, admin)).rows}
    assert after == before == {"A": D(5), "P": D(2)}
    cardex = await reports.cardex(db, admin, ids["A"])
    assert [(r[1], r[5], r[6]) for r in cardex.rows] == [
        ("موجودی اول دوره", D(7), None), ("حواله خروج", None, D(2))]

    async with db.session() as s:
        remaining = (await s.scalars(select(Document))).all()
        assert {(d.doc_type, d.fiscal_year) for d in remaining} == {
            (DocType.LOAN_OUT, 1404), (DocType.OPENING, 1405), (DocType.ISSUE, 1405)}
        assert await s.scalar(select(func.min(StockLedger.doc_date))) == dt.date(2026, 3, 21)
        assert await s.scalar(select(Item.id).where(Item.code == "C")) is None
        actions = set((await s.scalars(select(AuditLog.action))).all())
    assert "fiscal.year_closed" in actions

    # The carried loan can still be returned in the new year...
    await docs.create_and_post(db, admin, docs.DocumentInput(
        DocType.LOAN_RETURN, NEW, env["wh"], [docs.LineInput(ids["P"], env["u"]["عدد"], D(4))],
        person_id=env["emp"], related_document_id=env["loan"]))
    assert await docs.outstanding_loans(db, admin) == []
    # ...but nothing dated in the closed year can be created or cancelled.
    with pytest.raises(ValidationError, match="بسته"):
        await docs.create_document(db, admin, docs.DocumentInput(
            DocType.RECEIPT, Y1404, env["wh"], [docs.LineInput(ids["A"], env["u"]["عدد"], D(1))]))
    with pytest.raises(ValidationError, match="بسته"):
        await docs.cancel_document(db, admin, env["loan"])


async def test_archive_and_audit_export(env):
    db = env["db"]
    async with db.session() as s:  # an event from last year (audit times are real clock times)
        s.add(AuditLog(at=dt.datetime(2025, 10, 1, 9, 0), action="auth.login", machine="PC-1"))
    result = await _close(env)
    state = await load_state(db)
    assert state.closed_through == 1404
    [archive] = state.archive_list()
    assert archive.year == 1404 and Path(archive.database).exists()
    archived = fiscal.open_archive(db, None, "", archive)
    try:
        async with archived.session() as s:
            old = await s.scalar(select(func.count()).select_from(Document)
                                 .where(Document.fiscal_year == 1404))
        assert old == 3  # the archive still has the full year
        with pytest.raises(Exception):  # noqa: B017 - archive is read-only
            async with archived.session() as s:
                s.add(Person(code="X", name="x", name_normalized="x"))
    finally:
        await archived.dispose()
    wb = load_workbook(result.audit_file)
    assert result.audit_rows_exported == 1
    assert any("auth.login" in [c.value for c in row] for row in wb.active.iter_rows())
    async with db.session() as s:  # wiped from the active database
        assert await s.scalar(select(AuditLog.id).where(AuditLog.machine == "PC-1")) is None


async def test_fresh_start(env):
    db, admin = env["db"], env["admin"]
    await _close(env, CloseOptions.fresh_start())
    async with db.session() as s:
        codes = set((await s.scalars(select(Item.code))).all())
        persons = set((await s.scalars(select(Person.name))).all())
        statuses = set((await s.scalars(select(Document.status))).all())
    # Only what the new year still needs survives: the item issued in 1405, the loaned pallet
    # and the person holding it.
    assert codes == {"A", "P"} and persons == {"علی"}
    assert DocStatus.POSTED in statuses
    rows = (await reports.stock_balance(db, admin)).rows
    assert {r[0]: r[4] for r in rows} == {"A": D(-2)}  # new-year issue without carried stock


MARIADB_URL = os.environ.get("CASPIAN_TEST_MARIADB_URL")
MARIADB_ROOT = os.environ.get("CASPIAN_TEST_MARIADB_ROOT")


@pytest.mark.skipif(not (MARIADB_URL and MARIADB_ROOT), reason="MariaDB admin credentials not set")
async def test_year_end_on_mariadb(tmp_path):
    from sqlalchemy import text
    from sqlalchemy.engine import make_url

    from caspian.db.base import Base
    from caspian.db.bootstrap import prepare
    from caspian.db.database import Database, DbConfig
    from caspian.services import auth
    from caspian.services.backup import MariaDbDumper

    url = make_url(MARIADB_URL)
    config = DbConfig(url.host, url.port or 3306, url.database, url.username)
    root_user, root_password = MARIADB_ROOT.split(":", 1)
    db = Database(MARIADB_URL)
    archive_db = f"{url.database}_1404"
    try:
        async with db.engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)
            await conn.execute(text("DROP TABLE IF EXISTS alembic_version"))
        await prepare(db)
        admin = (await auth.login(db, "admin", "admin")).actor
        await auth.change_password(db, admin, "admin", "Str0ngPass")
        await auth.set_pin(db, admin, "Str0ngPass", "4826")
        unit = (await master.list_units(db))[0].id
        wh = (await master.list_warehouses(db))[0].id
        item = await items.create_item(db, admin, ItemInput("A", "دریل", unit))
        await docs.create_and_post(db, admin, docs.DocumentInput(
            DocType.RECEIPT, Y1404, wh, [docs.LineInput(item, unit, D(9))]))
        result = await fiscal.run_year_end(
            db, admin, 1404, CloseOptions(), await _approval(db, admin),
            MariaDbDumper(config, url.password), PASSWORD, tmp_path, config, root_user, root_password)
        assert result.archive == archive_db and result.opening_documents == 1
        [archive] = (await load_state(db)).archive_list()
        archived = fiscal.open_archive(db, config, url.password, archive)
        try:
            async with archived.session() as s:
                assert await s.scalar(select(func.count()).select_from(Document)) == 1
            with pytest.raises(Exception):  # noqa: B017 - SELECT-only grants + read-only session
                async with archived.session() as s:
                    s.add(Person(code="X", name="x", name_normalized="x"))
        finally:
            await archived.dispose()
        assert {r[0]: r[4] for r in (await reports.stock_balance(db, admin)).rows} == {"A": D(9)}
    finally:
        await db.dispose()
        root = Database(DbConfig(config.host, config.port, "", root_user).url(root_password, database=""))
        async with root.engine.connect() as conn:
            await conn.execute(text(f"DROP DATABASE IF EXISTS `{archive_db}`"))
        await root.dispose()
