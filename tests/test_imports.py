import io
from decimal import Decimal

import pytest
from openpyxl import Workbook

from caspian.core.text import normalize
from caspian.db.models import (
    BatchStatus,
    DocStatus,
    DocType,
    ImportKind,
    ImportSource,
    LineStatus,
    Resolution,
)
from caspian.services import documents as docs
from caspian.services import imports, items, master, protected
from caspian.services.errors import ApprovalError, PermissionDenied, ValidationError
from caspian.services.import_files import (
    RawRow,
    detect_mapping,
    read_csv,
    read_excel,
    read_word,
    rows_from_table,
)
from caspian.services.items import ItemInput
from caspian.services.protected import ProtectedAction

# ----- file readers -----


def _xlsx(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_excel_and_mapping():
    table = read_excel(_xlsx([
        ["كد", "شرح کالا", "تعداد", "واحد", "فی", "توضیح"],
        [1001, "دریل", 5, "عدد", 1500000.0, "x"],
        [None, None, None, None, None, None],
        ["A-2", "میز", "۳", "", "", ""],
    ]), "list.xlsx")
    assert table.headers[:3] == ["كد", "شرح کالا", "تعداد"]
    mapping = detect_mapping(table.headers)
    assert mapping == {"code": 0, "name": 1, "qty": 2, "unit_name": 3, "unit_price": 4}
    rows = rows_from_table(table, mapping)
    assert len(rows) == 2  # blank row dropped
    assert (rows[0].code, rows[0].qty, rows[0].unit_price) == ("1001", Decimal(5), Decimal(1500000))
    assert rows[1].qty == Decimal(3)  # Persian digits


def test_read_csv_encodings_and_delimiters():
    # Windows-1256 has only the Arabic yeh; normalization makes it match Persian names.
    text = "name;qty\nدريل;2\nميز;abc\n"
    for encoding in ("utf-8-sig", "cp1256"):
        table = read_csv(text.encode(encoding))
        rows = rows_from_table(table, detect_mapping(table.headers))
        assert normalize(rows[0].name) == normalize("دریل") and rows[0].qty == Decimal(2)
        assert "عدد نیست" in rows[1].error


def test_read_word_uses_largest_table():
    from docx import Document

    d = Document()
    d.add_paragraph("لیست خرید")
    small = d.add_table(rows=1, cols=2)
    small.rows[0].cells[0].text = "x"
    big = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("نام کالا", "تعداد"), ("پیچ", "100"), ("لولا", "20")]):
        big.rows[r].cells[0].text, big.rows[r].cells[1].text = a, b
    buf = io.BytesIO()
    d.save(buf)
    table = read_word(buf.getvalue())
    assert table.headers == ["نام کالا", "تعداد"] and len(table.rows) == 2


def test_bad_files():
    with pytest.raises(ValidationError):
        read_excel(b"not an xlsx")
    with pytest.raises(ValidationError):
        rows_from_table(read_csv("qty,unit\n1,عدد\n".encode()), {"qty": 0})


# ----- matching & batches -----


@pytest.fixture
async def catalog(db, admin):
    u = {x.name: x.id for x in await master.list_units(db)}
    drill = await items.create_item(db, admin, ItemInput(
        "1001", "دریل شارژی بوش", u["عدد"], units=[(u["جعبه"], Decimal(4))],
        barcodes=[("626111", None)]))
    screw = await items.create_item(db, admin, ItemInput("1002", "پیچ ام دی اف ۴ در ۴۰", u["عدد"]))
    wh = (await master.list_warehouses(db))[0].id
    return {"u": u, "drill": drill, "screw": screw, "wh": wh}


async def _lines(db, admin, batch_id):
    return (await imports.get_batch(db, admin, batch_id)).lines


async def test_item_import_statuses(db, admin, catalog):
    rows = [
        RawRow(code="1001", name="دریل شارژی بوش"),  # exact code, same name
        RawRow(code="1001", name="x"),  # duplicate within file
        RawRow(code="1002", name="صندلی گردان"),  # code exists, very different name
        RawRow(code="9001", name="پیچ ام‌دی‌اف ۴ در ۴۰ میل"),  # likely duplicate by name
        RawRow(code="9002", name="کمد فلزی"),  # new
        RawRow(code="9003", name="میز", error="«x» در ستون مقدار عدد نیست"),
        RawRow(),  # empty
    ]
    batch = await imports.create_batch(db, admin, ImportKind.ITEMS, ImportSource.EXCEL, rows, "t")
    lines = await _lines(db, admin, batch)
    assert [ln.status for ln in lines] == [
        LineStatus.EXISTING_MATCH, LineStatus.CONFLICT, LineStatus.CONFLICT, LineStatus.CONFLICT,
        LineStatus.NEW, LineStatus.ERROR, LineStatus.IGNORED]
    assert "ردیف 1" in lines[1].message
    assert lines[3].match_item_id == catalog["screw"] and lines[3].candidates
    assert [ln.resolution for ln in lines] == [
        Resolution.MATCH, None, None, None, Resolution.CREATE, None, Resolution.IGNORE]
    with pytest.raises(ValidationError, match="نیاز به تصمیم"):
        await imports.apply_batch(db, admin, batch)

    for ln in (lines[1], lines[2], lines[5]):
        await imports.set_resolution(db, admin, ln.id, Resolution.IGNORE)
    await imports.set_resolution(db, admin, lines[3].id, Resolution.MATCH, catalog["screw"])
    result = await imports.apply_batch(db, admin, batch)
    assert (result.created_items, result.updated_items, result.document_id) == (1, 0, None)
    assert [r.name for r in await items.search_items(db, admin, "کمد")] == ["کمد فلزی"]
    with pytest.raises(ValidationError):  # already applied
        await imports.apply_batch(db, admin, batch)


async def test_overwrite_requires_pin(db, admin, catalog):
    batch = await imports.create_batch(db, admin, ImportKind.ITEMS, ImportSource.CSV, [
        RawRow(code="1001", name="دریل شارژی بوش مدل جدید", barcode="626999",
               reorder_point=Decimal(3), category_name="ابزار برقی")])
    [line] = await _lines(db, admin, batch)
    await imports.set_resolution(db, admin, line.id, Resolution.OVERWRITE)
    assert imports.needs_overwrite_approval(await imports.get_batch(db, admin, batch))
    with pytest.raises(ApprovalError):
        await imports.apply_batch(db, admin, batch)
    approval = await protected.approve(db, admin, ProtectedAction.IMPORT_OVERWRITE, "admin", "4826")
    result = await imports.apply_batch(db, admin, batch, approval)
    assert result.updated_items == 1
    detail = await items.get_item(db, admin, catalog["drill"])
    assert detail.input.name == "دریل شارژی بوش مدل جدید"
    assert detail.input.reorder_point == Decimal(3)
    assert {b for b, _ in detail.input.barcodes} == {"626111", "626999"}
    assert detail.input.category_id is not None


async def test_stock_import_creates_draft_document(db, admin, catalog):
    rows = [
        RawRow(barcode="626111", qty=Decimal(2), unit_name="جعبه"),  # by barcode, in boxes
        RawRow(name="دریل شارژي بوش", qty=Decimal(1)),  # exact name (Arabic yeh)
        RawRow(name="پیچ ام دی اف ۴ در ۴۰ میلی", qty=Decimal(100)),  # close fuzzy
        RawRow(barcode="777000", qty=Decimal(3)),  # unknown barcode -> new, needs a name
        RawRow(name="دریل", qty=Decimal(0)),  # bad qty
        RawRow(code="1001", qty=Decimal(1), unit_name="کارتن"),  # unit not defined for item
    ]
    batch = await imports.create_batch(
        db, admin, ImportKind.STOCK, ImportSource.SCAN, rows, "اسکن", DocType.RECEIPT, catalog["wh"])
    lines = await _lines(db, admin, batch)
    assert [ln.status for ln in lines] == [
        LineStatus.EXISTING_MATCH, LineStatus.EXISTING_MATCH, LineStatus.EXISTING_MATCH,
        LineStatus.NEW, LineStatus.ERROR, LineStatus.ERROR]
    assert "تقریبی" in lines[2].message
    assert Resolution.CREATE not in lines[3].allowed  # no name yet
    assert lines[3].resolution is None

    await imports.update_line(db, admin, lines[3].id, name="اره عمود بر")
    await imports.update_line(db, admin, lines[4].id, qty=Decimal(2))
    await imports.set_resolution(db, admin, lines[5].id, Resolution.IGNORE)
    lines = await _lines(db, admin, batch)
    assert lines[3].resolution == Resolution.CREATE and lines[4].resolution == Resolution.MATCH

    result = await imports.apply_batch(db, admin, batch)
    assert result.created_items == 1
    detail = await docs.get_document(db, admin, result.document_id)
    assert detail.status is DocStatus.DRAFT  # never posted automatically
    assert [(ln.item_name, ln.unit_name, ln.base_qty) for ln in detail.lines] == [
        ("دریل شارژی بوش", "جعبه", Decimal(8)), ("دریل شارژی بوش", "عدد", Decimal(1)),
        ("پیچ ام دی اف ۴ در ۴۰", "عدد", Decimal(100)), ("اره عمود بر", "عدد", Decimal(3)),
        ("دریل شارژی بوش", "عدد", Decimal(2))]
    assert (await imports.list_batches(db, admin))== []
    assert (await imports.list_batches(db, admin, BatchStatus.APPLIED))[0].id == batch


async def test_ai_can_draft_but_not_apply(db, admin, catalog):
    ai = admin.as_ai()
    batch = await imports.create_batch(db, ai, ImportKind.STOCK, ImportSource.TEXT,
                                       [RawRow(name="دریل شارژی بوش", qty=Decimal(5))],
                                       "پیام دستیار", DocType.RECEIPT, catalog["wh"])
    with pytest.raises(PermissionDenied):
        await imports.apply_batch(db, ai, batch)
    await imports.apply_batch(db, admin, batch)  # the human applies it


async def test_discard(db, admin, catalog):
    batch = await imports.create_batch(db, admin, ImportKind.ITEMS, ImportSource.EXCEL,
                                       [RawRow(name="x")])
    await imports.discard_batch(db, admin, batch)
    assert await imports.list_batches(db, admin) == []
    with pytest.raises(ValidationError):
        await imports.update_line(db, admin, (await _lines(db, admin, batch))[0].id, name="y")


async def test_deleting_import_draft_reopens_batch(db, admin, catalog):
    """Deleting the draft an import created failed with FK error 1451 (#3). Now the batch goes
    back to OPEN and re-applying it reuses the items created the first time (no duplicates)."""
    rows = [RawRow(code="1001", qty=Decimal(2)), RawRow(barcode="777000", name="اره عمود بر", qty=Decimal(3))]
    batch = await imports.create_batch(
        db, admin, ImportKind.STOCK, ImportSource.SCAN, rows, "اسکن", DocType.RECEIPT, catalog["wh"])
    first = await imports.apply_batch(db, admin, batch)
    assert first.created_items == 1

    await docs.delete_draft(db, admin, first.document_id)

    detail = await imports.get_batch(db, admin, batch)
    assert detail.row.status is BatchStatus.OPEN and detail.result_document_id is None
    assert [ln.status for ln in detail.lines] == [LineStatus.EXISTING_MATCH] * 2
    second = await imports.apply_batch(db, admin, batch)
    assert second.created_items == 0
    redo = await docs.get_document(db, admin, second.document_id)
    assert redo.status is DocStatus.DRAFT and len(redo.lines) == 2
    assert len(await items.search_items(db, admin, "اره عمود بر")) == 1


# ----- manual entry and similar-item suggestions (#13, #14) -----


async def test_manual_rows_become_a_reviewable_batch(db, admin, catalog):
    batch = await imports.create_batch(
        db, admin, ImportKind.STOCK, ImportSource.MANUAL,
        [RawRow(code="1001", qty=Decimal(2)), RawRow(name="قفسه فلزی", qty=Decimal(1))],
        "ورود دستی", DocType.RECEIPT, catalog["wh"])
    detail = await imports.get_batch(db, admin, batch)
    assert detail.row.source is ImportSource.MANUAL
    assert imports.SOURCE_NAMES[ImportSource.MANUAL] == "ورود دستی"
    assert [ln.status for ln in detail.lines] == [LineStatus.EXISTING_MATCH, LineStatus.NEW]


async def test_matched_rows_show_the_item_code_and_alternatives(db, admin, catalog):
    batch = await imports.create_batch(
        db, admin, ImportKind.STOCK, ImportSource.EXCEL,
        [RawRow(name="پیچ ام دی اف ۴ در ۴۰ میلی", qty=Decimal(100)),
         RawRow(barcode="626111", qty=Decimal(1))],
        "فاکتور", DocType.RECEIPT, catalog["wh"])
    approx, by_barcode = await _lines(db, admin, batch)
    assert approx.status is LineStatus.EXISTING_MATCH and approx.code == ""
    assert approx.match_code == "1002"  # the code column is no longer empty for matched rows
    assert approx.candidates and approx.candidates[0][0] == catalog["screw"]  # alternatives kept
    assert by_barcode.match_code == "1001"


async def test_new_variant_of_a_similar_item(db, admin, catalog):
    """«نسخه جدید از همین کالا»: a new item named after the similar one, name editable (#14)."""
    batch = await imports.create_batch(
        db, admin, ImportKind.STOCK, ImportSource.EXCEL,
        [RawRow(name="دریل شارژی بوش مدل ۲", qty=Decimal(1))], "فاکتور", DocType.RECEIPT, catalog["wh"])
    [line] = await _lines(db, admin, batch)
    # Matched (approximately) or in conflict: either way the similar item is offered.
    assert line.status in (LineStatus.EXISTING_MATCH, LineStatus.CONFLICT)
    assert line.candidates[0][0] == catalog["drill"]
    assert imports.variant_name("دریل شارژی بوش") == "دریل شارژی بوش - "
    await imports.create_variant(db, admin, line.id, "دریل شارژی بوش - مدل ۲")
    [line] = await _lines(db, admin, batch)
    assert line.resolution is Resolution.CREATE and line.name == "دریل شارژی بوش - مدل ۲"
    result = await imports.apply_batch(db, admin, batch)
    assert result.created_items == 1
    assert [r.name for r in await items.search_items(db, admin, "مدل ۲")] == ["دریل شارژی بوش - مدل ۲"]
    with pytest.raises(ValidationError):
        await imports.create_variant(db, admin, line.id, "   ")
