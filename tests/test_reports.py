import datetime as dt
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from caspian.db.models import DocType, PersonKind
from caspian.services import auth, items, master, reports, users
from caspian.services import documents as docs
from caspian.services.errors import PermissionDenied
from caspian.services.excel_export import xlsx_bytes
from caspian.services.items import ItemInput
from caspian.services.reports import ReorderParams, compute_burn_rate

D = Decimal
TODAY = dt.date(2026, 9, 27)


@pytest.fixture
async def env(db, admin):
    u = {x.name: x.id for x in await master.list_units(db)}
    wh = (await master.list_warehouses(db))[0].id
    wh2 = await master.save_warehouse(db, admin, "02", "دوم")
    drill = await items.create_item(db, admin, ItemInput(
        "1001", "دریل", u["عدد"], units=[(u["جعبه"], D(4))], reorder_point=D(5)))
    table = await items.create_item(db, admin, ItemInput("2001", "میز", u["عدد"]))

    async def post(doc_type, date, wh_id, *lines, **kw):
        return await docs.create_and_post(db, admin, docs.DocumentInput(
            doc_type, date, wh_id, [docs.LineInput(*ln) for ln in lines], **kw))

    # 2 boxes at 4,000,000 per box -> 1,000,000 per piece
    await post(DocType.RECEIPT, TODAY - dt.timedelta(40), wh, (drill, u["جعبه"], D(2), D(4_000_000)))
    await post(DocType.RECEIPT, TODAY - dt.timedelta(30), wh2, (drill, u["عدد"], D(2), D(1_100_000)))
    await post(DocType.ISSUE, TODAY - dt.timedelta(20), wh, (drill, u["عدد"], D(3)))
    cancelled = await post(DocType.ISSUE, TODAY - dt.timedelta(10), wh, (drill, u["عدد"], D(1)))
    await docs.cancel_document(db, admin, cancelled)
    await post(DocType.RECEIPT, TODAY - dt.timedelta(5), wh, (table, u["عدد"], D(1)))
    return {"u": u, "wh": wh, "wh2": wh2, "drill": drill, "table": table}


async def test_stock_balance_with_last_price(db, admin, env):
    report = await reports.stock_balance(db, admin)
    rows = {r[0]: r for r in report.rows}
    assert rows["1001"][4] == D(7)  # 8 + 2 - 3
    assert rows["1001"][6] == D(1_100_000)  # most recent receipt price, per piece
    assert rows["1001"][7] == D(7_700_000)
    assert rows["2001"][6] is None and rows["2001"][7] is None
    assert report.totals[7] == D(7_700_000)
    only_wh2 = await reports.stock_balance(db, admin, env["wh2"])
    assert [(r[0], r[4]) for r in only_wh2.rows] == [("1001", D(2))]


async def test_cardex_running_balance_hides_cancelled(db, admin, env):
    report = await reports.cardex(db, admin, env["drill"], env["wh"])
    assert [(r[5], r[6], r[7]) for r in report.rows] == [(D(8), None, D(8)), (None, D(3), D(5))]
    later = await reports.cardex(db, admin, env["drill"], date_from=TODAY - dt.timedelta(35))
    assert later.rows[0][1] == "مانده از قبل" and later.rows[0][7] == D(8)
    assert later.totals == {5: D(2), 6: D(3), 7: D(7)}


def test_burn_rate_math():
    params = ReorderParams(lookback_days=90, lead_time_days=14, cover_days=60, safety_days=7)
    r = compute_burn_rate(1, "1", "پیچ", "عدد", on_hand=D(30), consumed=D(180), reorder_point=None,
                          params=params)
    assert r.daily == D(2)
    assert r.coverage_days == 15
    assert r.suggested_reorder_point == D(42)  # 2 * (14 + 7)
    assert r.suggested_order_qty == D(132)  # 2 * (14 + 60 + 7) - 30
    assert "۱۳۲ عدد" in r.suggestion_text and "۲٫۲ ماه" in r.suggestion_text
    idle = compute_burn_rate(1, "1", "x", "عدد", D(5), D(0), None, params)
    assert idle.coverage_days is None and idle.suggested_order_qty == 0
    assert idle.suggestion_text == "نیازی به سفارش نیست."


async def test_burn_rates_from_issues(db, admin, env):
    params = ReorderParams(lookback_days=30, lead_time_days=10, cover_days=30, safety_days=0)
    rates = {r.code: r for r in await reports.burn_rates(db, admin, params, today=TODAY)}
    drill = rates["1001"]
    assert drill.consumed == D(3)  # cancelled issue excluded
    assert drill.daily == D("0.1")
    assert drill.suggested_order_qty == D(0)  # 7 on hand >= 0.1 * 40
    params = ReorderParams(lookback_days=30, lead_time_days=60, cover_days=60, safety_days=0)
    needing = await reports.burn_rates(db, admin, params, only_needing_order=True, today=TODAY)
    assert [r.code for r in needing] == []  # 7 on hand > reorder point 5
    await items.set_reorder_points(db, admin, {env["drill"]: D(8)})
    needing = await reports.burn_rates(db, admin, params, only_needing_order=True, today=TODAY)
    assert [(r.code, r.suggested_order_qty) for r in needing] == [("1001", D(5))]


async def test_loans_and_activity_reports(db, admin, env):
    await users.create_user(db, admin, "neda", "", "Neda#2026", "viewer")
    viewer = (await auth.login(db, "neda", "Neda#2026")).actor
    with pytest.raises(PermissionDenied):
        await reports.user_activity(db, viewer)
    activity = await reports.user_activity(db, admin, admin.user_id)
    actions = [r[2] for r in activity.rows]
    assert "ثبت نهایی سند" in actions and "ابطال سند" in actions
    pallet = await items.create_item(db, admin, ItemInput("3001", "پالت", env["u"]["عدد"],
                                                          is_returnable=True))
    person = await master.save_person(db, admin, "علی", PersonKind.EMPLOYEE)
    await docs.create_and_post(db, admin, docs.DocumentInput(
        DocType.RECEIPT, TODAY, env["wh"], [docs.LineInput(pallet, env["u"]["عدد"], D(3))]))
    await docs.create_and_post(db, admin, docs.DocumentInput(
        DocType.LOAN_OUT, TODAY, env["wh"], [docs.LineInput(pallet, env["u"]["عدد"], D(2))],
        person_id=person))
    loans = await reports.loans_report(db, admin)
    assert [(r[2], r[4], r[5]) for r in loans.rows] == [("علی", "پالت", D(2))]


async def test_counter_cannot_see_stock_reports(db, admin, env):
    await users.create_user(db, admin, "shomar", "", "Count#2026", "counter")
    counter = (await auth.login(db, "shomar", "Count#2026")).actor
    with pytest.raises(PermissionDenied):
        await reports.stock_balance(db, counter)


async def test_excel_export_is_rtl_and_numeric(db, admin, env):
    report = await reports.stock_balance(db, admin, include_zero=True)
    wb = load_workbook(io.BytesIO(xlsx_bytes(report)))
    ws = wb.active
    assert ws.sheet_view.rightToLeft
    assert ws.cell(1, 1).value == "گزارش موجودی کالا"
    values = [[c.value for c in row] for row in ws.iter_rows()]
    header = next(i for i, r in enumerate(values) if r[0] == "کد")
    first = values[header + 1]
    assert first[0] == "1001" and first[4] == 7 and first[7] == 7_700_000
    assert values[-1][0] == "جمع"
    cardex = await reports.cardex(db, admin, env["drill"])
    ws = load_workbook(io.BytesIO(xlsx_bytes(cardex))).active
    dates = [r[0].value for r in ws.iter_rows() if r[0].value and str(r[0].value).startswith("1405/")]
    assert dates  # Jalali dates written as text
