"""Phase 2.2: Persian names, mobile numbers and passwords are cleaned and checked exactly like Caspian
Server does, and search finds people/items/places whatever ی/ي, ک/ك, ZWNJ or digits were typed."""

import os

import pytest

from caspian.core.persian import (
    NAME_ERROR,
    clean_person_name,
    name_problem,
    normalize_mobile,
    password_problem,
    password_rules,
)
from caspian.core.text import normalize, normalize_location


@pytest.mark.parametrize(("raw", "clean"), [
    ("عَلي", "علی"),
    ("  رضا   كريمي ", "رضا کریمی"),
    ("مؤمنی", "مؤمنی"),
    ("رئیسی", "رئیسی"),
    ("علاءالدین", "علاءالدین"),
    ("اسماعیل‌زاده", "اسماعیل‌زاده"),
    ("‌اسماعیل‌‌زاده‌", "اسماعیل‌زاده"),
    ("مهـــدی", "مهدی"),
    ("‏نازنین‎", "نازنین"),
    ("علی ‌ رضا", "علی رضا"),
])
def test_names_are_cleaned_like_the_server(raw, clean):
    assert clean_person_name(raw) == clean
    assert name_problem(clean_person_name(raw)) is None


@pytest.mark.parametrize("bad", ["Ali", "علی۲", "ع", "علی😀", "x" * 3, "ا" * 51, "علی-رضا"])
def test_bad_names_get_the_servers_sentence(bad):
    assert name_problem(clean_person_name(bad)) == NAME_ERROR


@pytest.mark.parametrize("raw", ["09121234567", "9121234567", "+989121234567", "00989121234567",
                                 "۰۹۱۲ ۱۲۳ ۴۵۶۷", "٠٩١٢-١٢٣-٤٥٦٧", "+98 (912) 123-4567"])
def test_mobile_forms(raw):
    assert normalize_mobile(raw) == "09121234567"


@pytest.mark.parametrize("raw", ["", "0912123456", "08121234567", "091212345678", "admin"])
def test_not_a_mobile(raw):
    assert normalize_mobile(raw) is None


def test_password_rules_live_checklist():
    assert password_problem("anbar-1403") is None
    assert all(ok for _, ok in password_rules("anbar-1403"))
    assert password_problem("1234567") is not None  # too short
    assert password_problem("12345678") is not None  # common
    assert password_problem("aaaabbbb") is not None  # 2 different characters
    assert password_problem("رمز" * 13) is not None  # more than 72 bytes


@pytest.mark.parametrize(("typed", "stored"), [
    ("علي رضايي", "علی رضایی"),
    ("كيسه زباله", "کیسه زباله"),
    ("اسماعيل زاده", "اسماعیل‌زاده"),
    ("لامپ 20 وات", "لامپ ۲۰ وات"),
])
def test_search_ignores_arabic_letters_zwnj_and_digits(typed, stored):
    assert normalize(typed) == normalize(stored)


def test_location_search_ignores_digits_and_spaces():
    assert normalize_location("پلاک 68") == normalize_location("پلاک۶۸")


MARIADB_URL = os.environ.get("CASPIAN_TEST_MARIADB_URL")


@pytest.mark.skipif(not MARIADB_URL, reason="CASPIAN_TEST_MARIADB_URL not set")
async def test_persian_names_round_trip_in_mariadb_utf8mb4():
    """Database, tables and connection are utf8mb4: ZWNJ, hamza and Persian digits come back intact."""
    from sqlalchemy import select, text

    from caspian.db.base import Base
    from caspian.db.bootstrap import prepare
    from caspian.db.database import Database
    from caspian.db.models import User
    from caspian.services import auth

    db = Database(MARIADB_URL)
    try:
        async with db.engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)
            await conn.execute(text("DROP TABLE IF EXISTS alembic_version"))
        await prepare(db)
        name = "علاءالدین مؤمنی رئیسی اسماعیل‌زاده ۱۲۳"
        admin = (await auth.login(db, "admin", "admin")).actor
        async with db.session() as s:
            (await s.get(User, admin.user_id)).full_name = name
        async with db.session() as s:
            assert (await s.scalar(select(User.full_name).where(User.id == admin.user_id))) == name
            charset = await s.scalar(text("SELECT @@character_set_connection"))
            table = await s.scalar(text(
                "SELECT TABLE_COLLATION FROM information_schema.TABLES "
                "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = 'users'"))
        assert charset == "utf8mb4" and table.startswith("utf8mb4")
    finally:
        await db.dispose()


def test_excel_export_keeps_persian_names_exactly():
    import io

    from openpyxl import load_workbook

    from caspian.services.excel_export import write_xlsx
    from caspian.services.reports import Column, ReportTable

    name = "اسماعیل‌زاده — مؤمنی"
    buffer = io.BytesIO()
    write_xlsx(ReportTable("مصرف", [Column("درخواست‌کننده", "text")], [[name]]), buffer)
    sheet = load_workbook(io.BytesIO(buffer.getvalue())).active
    assert any(cell.value == name for row in sheet.iter_rows() for cell in row)
    assert sheet.sheet_view.rightToLeft
