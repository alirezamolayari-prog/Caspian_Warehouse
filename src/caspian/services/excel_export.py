"""Export a ReportTable to a right-to-left Excel workbook."""

import datetime as dt
import io
import re
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from caspian import APP_DISPLAY_NAME
from caspian.core import jalali
from caspian.services.reports import ReportTable

_FORMATS = {"qty": "#,##0.####", "money": "#,##0", "int": "0"}
_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="DDF1EE")
_FONT = "Vazirmatn"


def _sheet_title(title: str) -> str:
    return re.sub(r"[\[\]:*?/\\]", " ", title)[:31] or "Report"


def _cell_value(value, kind: str):
    if value is None:
        return None
    if kind == "date" and isinstance(value, dt.date):
        text = jalali.format_date(value, persian_digits=False)
        if isinstance(value, dt.datetime):
            text += value.strftime(" %H:%M")
        return text
    if isinstance(value, Decimal):
        return float(value) if kind != "int" else int(value)
    return value


def write_xlsx(report: ReportTable, target: str | Path | io.BytesIO) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_title(report.title)
    ws.sheet_view.rightToLeft = True
    ncols = len(report.columns)

    ws.append([report.title])
    ws.cell(1, 1).font = Font(name=_FONT, bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 1))
    meta = [*report.meta, f"{APP_DISPLAY_NAME} — تهیه‌شده در "
            f"{jalali.format_date(dt.date.today(), persian_digits=False)}"]
    for line in meta:
        ws.append([line])
        ws.cell(ws.max_row, 1).font = Font(name=_FONT, color="555555")
    ws.append([])

    header_row = ws.max_row + 1
    ws.append([c.title for c in report.columns])
    for col in range(1, ncols + 1):
        cell = ws.cell(header_row, col)
        cell.font = Font(name=_FONT, bold=True)
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in report.rows:
        ws.append([_cell_value(v, c.kind) for v, c in zip(row, report.columns, strict=False)])
        r = ws.max_row
        for col, column in enumerate(report.columns, start=1):
            cell = ws.cell(r, col)
            cell.font = Font(name=_FONT)
            cell.border = _BORDER
            if column.kind in _FORMATS:
                cell.number_format = _FORMATS[column.kind]

    if report.totals:
        ws.append([])
        r = ws.max_row + 1
        ws.cell(r, 1, "جمع").font = Font(name=_FONT, bold=True)
        for index, value in report.totals.items():
            cell = ws.cell(r, index + 1, _cell_value(value, report.columns[index].kind))
            cell.font = Font(name=_FONT, bold=True)
            cell.number_format = _FORMATS.get(report.columns[index].kind, "General")

    for col, column in enumerate(report.columns, start=1):
        width = max([len(column.title), *(len(str(row[col - 1] or "")) for row in report.rows[:500])])
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 4, 10), 60)
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    if report.rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{header_row + len(report.rows)}"
    wb.save(target)


def xlsx_bytes(report: ReportTable) -> bytes:
    buffer = io.BytesIO()
    write_xlsx(report, buffer)
    return buffer.getvalue()
