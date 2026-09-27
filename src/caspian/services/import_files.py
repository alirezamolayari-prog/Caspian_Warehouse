"""Read tabular data from Excel / CSV / Word and map columns to import fields."""

import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from caspian.core.numbers import parse_decimal
from caspian.core.text import normalize
from caspian.db.models import ImportSource
from caspian.services.errors import ValidationError

MAX_ROWS = 20_000

# Import field -> (Persian label, header synonyms). Headers are compared normalized.
FIELDS: dict[str, tuple[str, tuple[str, ...]]] = {
    "code": ("کد کالا", ("کد", "کد کالا", "code", "sku", "item code", "شماره کالا")),
    "name": ("نام کالا", ("نام", "نام کالا", "شرح", "شرح کالا", "عنوان", "کالا", "name",
                          "item", "description", "product")),
    "barcode": ("بارکد", ("بارکد", "barcode", "ean", "upc")),
    "qty": ("مقدار", ("تعداد", "مقدار", "qty", "quantity", "count", "موجودی")),
    "unit_name": ("واحد", ("واحد", "unit", "uom", "واحد شمارش")),
    "unit_price": ("فی", ("فی", "قیمت", "قیمت واحد", "price", "unit price", "مبلغ واحد")),
    "category_name": ("گروه", ("گروه", "گروه کالا", "دسته", "دسته بندی", "category", "group")),
    "reorder_point": ("نقطه سفارش", ("نقطه سفارش", "حداقل موجودی", "reorder", "reorder point",
                                     "min", "min stock")),
}
NUMERIC_FIELDS = {"qty", "unit_price", "reorder_point"}


@dataclass
class TableData:
    headers: list[str]
    rows: list[list[str]]
    source: ImportSource
    file_name: str = ""


@dataclass
class RawRow:
    code: str = ""
    name: str = ""
    barcode: str = ""
    qty: Decimal | None = None
    unit_name: str = ""
    unit_price: Decimal | None = None
    category_name: str = ""
    reorder_point: Decimal | None = None
    raw: dict = field(default_factory=dict)
    error: str = ""  # parse problem found before matching


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores codes like 1001 as 1001.0
    return str(value).strip()


def _trim(headers: list[str], rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    rows = [r for r in rows if any(c for c in r)]
    width = max([len(headers), *(len(r) for r in rows)], default=0)
    headers = headers + [""] * (width - len(headers))
    rows = [r + [""] * (width - len(r)) for r in rows]
    return [h or f"ستون {i + 1}" for i, h in enumerate(headers)], rows[:MAX_ROWS]


def read_excel(data: bytes, file_name: str = "") -> TableData:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # corrupt or not an xlsx
        raise ValidationError("فایل اکسل قابل خواندن نیست (فقط .xlsx پشتیبانی می‌شود).") from exc
    ws = wb.active
    values = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    values = [r for r in values if any(r)]
    if not values:
        raise ValidationError("فایل اکسل خالی است.")
    headers, rows = _trim(values[0], values[1:])
    return TableData(headers, rows, ImportSource.EXCEL, file_name)


def read_csv(data: bytes, file_name: str = "") -> TableData:
    for encoding in ("utf-8-sig", "cp1256"):  # Excel-for-Persian often saves as cp1256
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValidationError("کدگذاری فایل CSV شناخته نشد (UTF-8 یا Windows-1256).")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    values = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    values = [r for r in values if any(r)]
    if not values:
        raise ValidationError("فایل CSV خالی است.")
    headers, rows = _trim(values[0], values[1:])
    return TableData(headers, rows, ImportSource.CSV, file_name)


def read_word(data: bytes, file_name: str = "") -> TableData:
    """Uses the largest table in the document (typical for lists converted from photos)."""
    from docx import Document as DocxDocument

    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ValidationError("فایل Word قابل خواندن نیست (فقط .docx پشتیبانی می‌شود).") from exc
    tables = [[[c.text.strip() for c in row.cells] for row in t.rows] for t in document.tables]
    tables = [[r for r in t if any(r)] for t in tables]
    tables = [t for t in tables if len(t) >= 2]
    if not tables:
        raise ValidationError("در فایل Word جدولی با حداقل یک ردیف داده پیدا نشد.")
    best = max(tables, key=len)
    headers, rows = _trim(best[0], best[1:])
    return TableData(headers, rows, ImportSource.WORD, file_name)


def read_file(path: str | Path) -> TableData:
    path = Path(path)
    data = path.read_bytes()
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return read_excel(data, path.name)
    if suffix in (".csv", ".txt"):
        return read_csv(data, path.name)
    if suffix == ".docx":
        return read_word(data, path.name)
    raise ValidationError("نوع فایل پشتیبانی نمی‌شود. از .xlsx، .csv یا .docx استفاده کنید.")


def detect_mapping(headers: list[str]) -> dict[str, int]:
    """Guess field -> column index from header names (Persian or English)."""
    mapping: dict[str, int] = {}
    normalized = [normalize(h) for h in headers]
    for field_name, (_label, synonyms) in FIELDS.items():
        wanted = {normalize(s) for s in synonyms}
        for i, header in enumerate(normalized):
            if header in wanted and i not in mapping.values():
                mapping[field_name] = i
                break
    return mapping


def rows_from_table(table: TableData, mapping: dict[str, int]) -> list[RawRow]:
    if "name" not in mapping and "code" not in mapping and "barcode" not in mapping:
        raise ValidationError("حداقل یکی از ستون‌های نام، کد یا بارکد کالا را مشخص کنید.")
    out = []
    for cells in table.rows:
        row = RawRow(raw=dict(zip(table.headers, cells, strict=False)))
        errors = []
        for field_name, index in mapping.items():
            value = cells[index] if index < len(cells) else ""
            if field_name in NUMERIC_FIELDS:
                try:
                    setattr(row, field_name, parse_decimal(value))
                except ValueError:
                    errors.append(f"«{value}» در ستون {FIELDS[field_name][0]} عدد نیست")
            else:
                setattr(row, field_name, value)
        row.error = "؛ ".join(errors)
        out.append(row)
    return out
